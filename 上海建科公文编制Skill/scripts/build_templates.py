#!/usr/bin/env python3
"""Build reusable Shanghai Jianke official-document DOCX templates."""

from pathlib import Path
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "assets" / "templates"
HEADER_IMAGE = ROOT / "assets" / "layout-references" / "normal-header-reference.png"


def set_east_asia(run, font):
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def set_table_borders(table, color="000000", size="8"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        node = borders.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:color"), color)


def set_cell_width(cell, width_mm):
    width = Mm(width_mm)
    cell.width = width
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width.twips))
    tc_w.set(qn("w:type"), "dxa")


def add_page_field(paragraph):
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "18")
    r_pr.append(size)
    run.append(r_pr)
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    field.append(run)
    paragraph._p.append(field)


def set_paragraph_bottom_border(paragraph, color="666666", size="6", space="2"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def configure_document(doc, with_header=True):
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(24)
    section.bottom_margin = Mm(22)
    section.left_margin = Mm(26)
    section.right_margin = Mm(24)
    section.header_distance = Mm(8)
    section.footer_distance = Mm(10)

    normal = doc.styles["Normal"]
    normal.font.size = Pt(12)
    normal.font.name = "宋体"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    normal.paragraph_format.line_spacing = Pt(25)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    if with_header and HEADER_IMAGE.exists():
        p = section.header.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        p.add_run().add_picture(str(HEADER_IMAGE), width=Mm(150))
        set_paragraph_bottom_border(p)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_field(footer)


def add_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(14)
    p.paragraph_format.line_spacing = Pt(26)
    run = p.add_run(text)
    set_east_asia(run, "黑体")
    run.bold = True
    run.font.size = Pt(18)


def add_body(doc, text, bold=False, indent=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.first_line_indent = Mm(7.4) if indent else Mm(0)
    run = p.add_run(text)
    set_east_asia(run, "宋体")
    run.font.size = Pt(12)
    run.bold = bold
    return p


def add_heading(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    set_east_asia(run, "黑体")
    run.bold = True
    run.font.size = Pt(12)


def add_signature(doc, lines):
    for text in lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.right_indent = Mm(8)
        p.paragraph_format.line_spacing = Pt(22)
        run = p.add_run(text)
        set_east_asia(run, "宋体")
        run.font.size = Pt(12)


def build_normal_notice(path):
    doc = Document()
    configure_document(doc)
    add_title(doc, "关于【事项】的通知")
    add_body(doc, "【主送机关】：", indent=False)
    add_body(doc, "【背景或依据】。为【目的】，公司决定【行动】，现将有关事项通知如下：")
    add_heading(doc, "一、【第一项】")
    add_body(doc, "【具体内容】")
    add_heading(doc, "二、【第二项】")
    add_body(doc, "【责任主体、工作要求、完成时限和提交方式】")
    add_body(doc, "附件：", indent=False)
    add_body(doc, "1、【附件名称】", indent=False)
    add_signature(doc, ["【发文主体/承办部门】", "【年】年【月】月【日】日"])
    doc.save(path)


def build_numbered_notice(path):
    doc = Document()
    configure_document(doc)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run("【内部文件编号】")
    set_east_asia(r, "宋体")
    r.font.size = Pt(8)
    add_title(doc, "关于【制度/权限事项】调整的通知")
    add_body(doc, "【主送机关】：", indent=False)
    add_body(doc, "为【管理目的】，根据【制度依据】，结合公司实际情况，现对【事项】进行调整，具体如下：")
    add_heading(doc, "一、【调整内容】")
    add_body(doc, "【条款或权限说明】")
    add_heading(doc, "二、【适用范围与生效】")
    add_body(doc, "本次调整适用于【范围】。自本通知印发之日起，【原制度或条款】同时废止。")
    add_body(doc, "特此通知。")
    add_signature(doc, ["上海建科工程咨询有限公司", "【年】年【月】月【日】日"])
    doc.save(path)


def build_party_notice(path):
    doc = Document()
    configure_document(doc, with_header=False)
    section = doc.sections[0]
    section.top_margin = Mm(22)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("中共上海建科工程咨询有限公司委员会")
    set_east_asia(r, "宋体")
    r.bold = True
    r.font.size = Pt(24)
    r.font.color.rgb = RGBColor(190, 36, 42)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("【党内发文字号】")
    set_east_asia(r, "宋体")
    r.font.size = Pt(10.5)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("──────────── ★ ────────────")
    r.font.color.rgb = RGBColor(190, 36, 42)
    r.font.size = Pt(10)

    add_title(doc, "关于【事项】的通知")
    add_body(doc, "【主送机关】：", indent=False)
    add_body(doc, "经【程序】、【审定机关】审定并【公示情况】，现将【事项】公布如下：")
    add_heading(doc, "一、【第一类】")
    add_body(doc, "【名单或内容】", indent=False)
    add_heading(doc, "二、【第二类】")
    add_body(doc, "【名单或内容】", indent=False)
    add_body(doc, "特此通知。")
    add_signature(doc, ["中共上海建科工程咨询有限公司委员会", "【年】年【月】月【日】日"])

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    r = p.add_run("【印发机关】　　　　　　　　　　　　　　【年】年【月】月【日】日印发")
    set_east_asia(r, "宋体")
    r.font.size = Pt(9)
    doc.save(path)


def build_form(path):
    doc = Document()
    configure_document(doc)
    add_title(doc, "【项目/事项】申报表")
    table = doc.add_table(rows=8, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    labels = [
        ("业务机构", "【填写】", "项目名称", "【填写】"),
        ("项目负责人", "【填写】", "项目人数", "【填写】"),
        ("开始日期", "【填写】", "计划完成日期", "【填写】"),
        ("当前阶段", "【填写】", "", ""),
        ("项目概况及特点", "【填写】", "", ""),
        ("工作难点", "【填写】", "", ""),
        ("预期成果及主要措施", "【填写】", "", ""),
        ("审核/推荐意见", "【签字及日期】", "", ""),
    ]
    for i, row in enumerate(table.rows):
        if i >= 4 and i <= 6:
            row.height = Mm(15)
            row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        elif i == 7:
            row.height = Mm(24)
            row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        widths = [32, 48, 32, 48]
        for j, cell in enumerate(row.cells):
            set_cell_width(cell, widths[j])
        if i >= 3:
            row.cells[1].merge(row.cells[3])
        for j, cell in enumerate(row.cells):
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if j < len(labels[i]):
                cell.text = labels[i][j]
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j in (0, 2) else WD_ALIGN_PARAGRAPH.LEFT
                for run in p.runs:
                    set_east_asia(run, "宋体")
                    run.font.size = Pt(10.5)
    doc.save(path)


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    build_normal_notice(OUT / "01-普通业务通知模板.docx")
    build_numbered_notice(OUT / "02-内部编号制度通知模板.docx")
    build_party_notice(OUT / "03-党委红头通知模板.docx")
    build_form(OUT / "04-申报评审表模板.docx")
    print(f"built templates in {OUT}")


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"template build failed: {exc}", file=sys.stderr)
        raise
